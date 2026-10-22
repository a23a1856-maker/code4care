"""Catalogue import (Excel/CSV): the new module requested for bulk medicine lists."""
from __future__ import annotations

import io
import json

import pytest
from openpyxl import Workbook, load_workbook

from app.seed_data import ean13
from tests.conftest import sign_in
from tests.helpers import count, csrf_token

HEADERS = [
    "barcode", "name", "brand", "generic", "strength", "form", "manufacturer",
    "condition", "uses", "dosage", "howToTake", "warnings", "sideEffects",
    "storage", "pack", "price", "stock", "expiry", "notes",
]

CODE_A = ean13("955001000901")
CODE_B = ean13("955001000902")
CODE_C = ean13("955001000903")


def sample(code: str = CODE_A, name: str = "Panadol Extra 500 mg", **overrides) -> dict:
    row = {
        "barcode": code,
        "name": name,
        "brand": "Panadol",
        "generic": "Paracetamol",
        "strength": "500 mg",
        "form": "Tablet",
        "manufacturer": "Haleon Malaysia",
        "condition": "Demam & Sakit",
        "uses": "Meredakan demam dan sakit ringan.",
        "dosage": "Dewasa: 1-2 tablet setiap 4-6 jam.",
        "howToTake": "Telan dengan segelas air.",
        "warnings": "Jangan melebihi 8 tablet dalam 24 jam.",
        "sideEffects": "Jarang: ruam.",
        "storage": "Simpan di bawah 30°C.",
        "pack": "1 kotak = 10 tablet",
        "price": "6.90",
        "stock": "24",
        "expiry": "2027-06",
        "notes": "Data demo - sahkan dengan label produk.",
    }
    row.update(overrides)
    return row


def workbook_bytes(rows: list[dict], headers: list[str] = HEADERS) -> bytes:
    workbook = Workbook()
    sheet = workbook.active
    sheet.append(headers)
    for row in rows:
        sheet.append([row.get(header, "") for header in headers])
    buffer = io.BytesIO()
    workbook.save(buffer)
    return buffer.getvalue()


def csv_bytes(rows: list[dict], headers: list[str]) -> bytes:
    lines = [",".join(headers)]
    for row in rows:
        cells = []
        for header in headers:
            value = str(row.get(header, "")).replace('"', '""')
            cells.append(f'"{value}"' if ("," in value or '"' in value) else value)
        lines.append(",".join(cells))
    return ("\n".join(lines) + "\n").encode("utf-8-sig")


def upload(client, signed, payload: bytes, filename: str = "catalog.xlsx", **extra):
    return client.post(
        "/catalog/import",
        data={"csrf_token": signed.csrf, **extra},
        files={"file": (filename, payload, "application/octet-stream")},
        headers={"HX-Request": "true"},
    )


def test_import_page_is_available_to_a_registered_user(client, world):
    signed = sign_in(client, world["users"]["owner_a"]["email"])
    page = client.get("/catalog/import")
    assert page.status_code == 200
    assert "Import Katalog" in page.text
    assert "Muat turun templat" in page.text
    assert "barcode" in page.text and "name" in page.text
    _ = signed


def test_import_xlsx_saves_rows_for_this_business_only(client, world):
    tenant_a = world["tenants"]["a"]
    tenant_b = world["tenants"]["b"]
    before_a = count(world["conn"], "medicines", tenant_a)
    before_b = count(world["conn"], "medicines", tenant_b)

    signed = sign_in(client, world["users"]["owner_a"]["email"])
    payload = workbook_bytes(
        [
            sample(CODE_A),
            sample(CODE_B, "Uphamol 250 Sirap", form="Sirap", strength="250 mg / 5 ml"),
            sample("C4C-0100", "Vitamin C 500 mg (katalog dalaman)", brand="Redoxon"),
        ]
    )
    response = upload(client, signed, payload)

    assert response.status_code == 200, response.text
    assert 'data-import="done"' in response.text
    assert "Import selesai" in response.text
    assert "3" in response.text and "ditambah" in response.text

    assert count(world["conn"], "medicines", tenant_a) == before_a + 3
    assert count(world["conn"], "medicines", tenant_b) == before_b

    stored = world["conn"].execute(
        "SELECT * FROM medicines WHERE tenant_id = ? AND barcode = ?", (tenant_a, CODE_A)
    ).fetchone()
    assert stored["name"] == "Panadol Extra 500 mg"
    assert stored["active_ingredient"] == "Paracetamol"
    assert stored["dosage_form"] == "Tablet"
    assert stored["manufacturer"] == "Haleon Malaysia"
    assert stored["category"] == "Demam & Sakit"
    assert stored["side_effects"] == "Jarang: ruam."
    assert stored["storage"] == "Simpan di bawah 30°C."
    assert stored["verified"] == 0, "imports are unverified unless the form says so"
    extra = json.loads(stored["extra_json"])
    assert extra["price"] == "6.90"
    assert extra["stock"] == "24"
    assert extra["expiry"] == "2027-06"
    assert extra["how_to_take"] == "Telan dengan segelas air."

    # The imported row is searchable, and the internal code works too.
    search = client.post(
        "/search", data={"q": CODE_A, "csrf_token": signed.csrf},
        headers={"HX-Request": "true"},
    )
    assert "Panadol Extra 500 mg" in search.text
    internal = client.post(
        "/search", data={"q": "C4C-0100", "csrf_token": signed.csrf},
        headers={"HX-Request": "true"},
    )
    assert "Vitamin C 500 mg" in internal.text


def test_reimport_updates_the_same_rows_without_duplicates(client, world):
    tenant_a = world["tenants"]["a"]
    signed = sign_in(client, world["users"]["owner_a"]["email"])

    first = upload(client, signed, workbook_bytes([sample(CODE_A), sample(CODE_B)]))
    assert first.status_code == 200
    after_first = count(world["conn"], "medicines", tenant_a)

    second = upload(
        client, signed,
        workbook_bytes([
            sample(CODE_A, "Panadol Extra (nama dikemas kini)", uses="Kegunaan dikemas kini."),
            sample(CODE_B),
        ]),
    )
    assert second.status_code == 200
    assert "dikemas kini" in second.text
    assert count(world["conn"], "medicines", tenant_a) == after_first

    row = world["conn"].execute(
        "SELECT name, uses FROM medicines WHERE tenant_id = ? AND barcode = ?",
        (tenant_a, CODE_A),
    ).fetchone()
    assert row["name"] == "Panadol Extra (nama dikemas kini)"
    assert row["uses"] == "Kegunaan dikemas kini."


def test_invalid_rows_are_reported_and_not_saved(client, world):
    tenant_a = world["tenants"]["a"]
    signed = sign_in(client, world["users"]["owner_a"]["email"])

    good = sample(CODE_A)
    missing_name = sample(CODE_B)
    missing_name["name"] = ""
    bad_code = sample("ABC")  # letters without a digit: not an EAN and not a SKU
    formula = sample(CODE_C, uses="=HYPERLINK(\"http://evil.example\",\"klik\")")
    duplicate = sample(CODE_A, "Pendua kod yang sama")

    response = upload(
        client, signed,
        workbook_bytes([good, missing_name, bad_code, formula, duplicate]),
    )
    assert response.status_code == 200, response.text
    assert "1" in response.text and "ditambah" in response.text
    assert "ditolak" in response.text

    # Row numbers in the report match the spreadsheet (header is row 1).
    assert ">3<" in response.text  # missing name
    assert ">4<" in response.text  # bad barcode
    assert ">5<" in response.text  # formula
    assert ">6<" in response.text  # duplicate barcode
    assert "formula tidak dibenarkan" in response.text
    assert "berulang" in response.text

    rows = world["conn"].execute(
        "SELECT barcode FROM medicines WHERE tenant_id = ? AND barcode IN (?, ?)",
        (tenant_a, CODE_B, CODE_C),
    ).fetchall()
    assert rows == []
    # Only the good row was written.
    assert world["conn"].execute(
        "SELECT COUNT(*) FROM medicines WHERE tenant_id = ? AND barcode = ?",
        (tenant_a, CODE_A),
    ).fetchone()[0] == 1


def test_dry_run_saves_nothing(client, world):
    tenant_a = world["tenants"]["a"]
    signed = sign_in(client, world["users"]["owner_a"]["email"])
    before = count(world["conn"], "medicines", tenant_a)

    response = upload(client, signed, workbook_bytes([sample(CODE_A)]), dry_run="1")

    assert response.status_code == 200
    assert 'data-import="preview"' in response.text
    assert "Pratonton" in response.text
    assert "tiada data disimpan" in response.text
    assert count(world["conn"], "medicines", tenant_a) == before


def test_verified_flag_controls_what_the_details_page_shows(client, world):
    tenant_a = world["tenants"]["a"]
    signed = sign_in(client, world["users"]["owner_a"]["email"])

    upload(client, signed, workbook_bytes([sample(CODE_A)]), verified="1", source="Katalog Jun 2026")
    row = world["conn"].execute(
        "SELECT id, verified, source, verified_at FROM medicines WHERE tenant_id = ? AND barcode = ?",
        (tenant_a, CODE_A),
    ).fetchone()
    assert row["verified"] == 1
    assert row["source"] == "Katalog Jun 2026"
    assert row["verified_at"]

    page = client.get(f"/medicines/{row['id']}")
    assert page.status_code == 200
    assert 'data-verified="true"' in page.text
    assert "Meredakan demam dan sakit ringan." in page.text
    assert "Jarang: ruam." in page.text
    assert "Katalog Jun 2026" in page.text

    # An unverified import shows the basics only, never the dosage or warnings.
    upload(client, signed, workbook_bytes([sample(CODE_B, "Ubat Tanpa Sah")]))
    plain = world["conn"].execute(
        "SELECT id FROM medicines WHERE tenant_id = ? AND barcode = ?", (tenant_a, CODE_B)
    ).fetchone()["id"]
    unverified_page = client.get(f"/medicines/{plain}")
    assert 'data-verified="false"' in unverified_page.text
    assert "belum disahkan" in unverified_page.text
    assert "Meredakan demam dan sakit ringan." not in unverified_page.text


def test_bad_files_are_refused_with_a_clear_message(client, world):
    tenant_a = world["tenants"]["a"]
    signed = sign_in(client, world["users"]["owner_a"]["email"])
    before = count(world["conn"], "medicines", tenant_a)

    not_excel = upload(client, signed, b"name,barcode\nA,1\n", "katalog.txt")
    assert not_excel.status_code == 422
    assert "Format fail tidak disokong" in not_excel.text

    broken = upload(client, signed, b"PK not really a workbook", "katalog.xlsx")
    assert broken.status_code == 422
    assert "tidak dapat dibaca" in broken.text

    too_big = upload(client, signed, b"x" * (5 * 1024 * 1024 + 10), "besar.xlsx")
    assert too_big.status_code == 422
    assert "terlalu besar" in too_big.text

    old_excel = upload(client, signed, b"whatever", "katalog.xls")
    assert old_excel.status_code == 422
    assert "simpan sebagai .xlsx" in old_excel.text

    missing_columns = upload(client, signed, workbook_bytes([{"foo": "bar"}], ["foo"]))
    assert missing_columns.status_code == 422
    assert "Kolum wajib tiada" in missing_columns.text

    assert count(world["conn"], "medicines", tenant_a) == before


def test_csv_with_the_sample_catalogue_columns(client, world):
    """The supplier file (catalog code4care.csv) shape imports as-is."""
    tenant_a = world["tenants"]["a"]
    signed = sign_in(client, world["users"]["owner_a"]["email"])
    headers = [
        "barcode", "name", "generic", "strength", "form", "condition", "regClass",
        "uses", "dosage", "howToTake", "warnings", "sideEffects", "storage",
        "regNo", "holder", "maker", "origin", "pack", "halal", "price", "stock",
        "expiry", "notes", "image",
    ]
    row = {
        "barcode": "C4C-0001",
        "name": "Panadol Actifast 500 mg",
        "generic": "Paracetamol",
        "strength": "500 mg",
        "form": "Tablet",
        "condition": "Demam & Sakit",
        "regClass": "OTC",
        "uses": "Mengurangkan demam dan melegakan sakit ringan.",
        "dosage": "Dewasa: 1-2 tablet setiap 4-6 jam.",
        "howToTake": "Telan dengan segelas air.",
        "warnings": "Jangan gabung dengan ubat lain yang mengandungi paracetamol.",
        "sideEffects": "Jarang: ruam, loya.",
        "storage": "Simpan di bawah 30°C.",
        "pack": "1 kotak = 10 tablet",
        "halal": "Tidak dinyatakan",
        "price": "6.9",
        "stock": "24",
        "expiry": "2027-06",
        "notes": "Data demo.",
        "image": "img/C4C-0001.jpg",
    }
    response = upload(client, signed, csv_bytes([row], headers), "catalog code4care.csv")

    assert response.status_code == 200, response.text
    assert "ditambah" in response.text

    stored = world["conn"].execute(
        "SELECT * FROM medicines WHERE tenant_id = ? AND barcode = ?", (tenant_a, "C4C-0001")
    ).fetchone()
    assert stored["name"] == "Panadol Actifast 500 mg"
    extra = json.loads(stored["extra_json"])
    assert extra["reg_class"] == "OTC"
    assert extra["price"] == "6.9"
    assert extra["image"] == "img/C4C-0001.jpg"


def test_template_download_is_a_real_workbook(client, world):
    signed = sign_in(client, world["users"]["owner_a"]["email"])
    response = client.get("/catalog/template.xlsx")
    assert response.status_code == 200
    assert "spreadsheetml" in response.headers["content-type"]
    assert "catalog-code4care-template.xlsx" in response.headers["content-disposition"]
    assert response.content[:2] == b"PK"

    workbook = load_workbook(io.BytesIO(response.content))
    sheet = workbook.active
    headers = [cell.value for cell in sheet[1]]
    assert headers[0] == "barcode"
    assert "name" in headers
    assert sheet.max_row == 3, "header plus two example rows"
    _ = signed


def test_import_requires_the_business_settings_permission(client, world):
    tenant_a = world["tenants"]["a"]
    before = count(world["conn"], "medicines", tenant_a)

    for role in ("pharmacist_a", "guest_a", "admin_a"):
        signed = sign_in(client, world["users"][role]["email"])
        page = client.get("/catalog/import", follow_redirects=False)
        assert page.status_code == 403, role
        posted = upload(client, signed, workbook_bytes([sample(CODE_A)]))
        assert posted.status_code == 403, role

    client.cookies.clear()
    assert client.get("/catalog/import", follow_redirects=False).status_code == 303
    assert client.post(
        "/catalog/import", data={}, files={"file": ("x.xlsx", b"x")}
    ).status_code == 401

    assert count(world["conn"], "medicines", tenant_a) == before


def test_import_needs_a_csrf_token(client, world):
    sign_in(client, world["users"]["owner_a"]["email"])
    response = client.post(
        "/catalog/import",
        files={"file": ("catalog.xlsx", workbook_bytes([sample(CODE_A)]))},
    )
    assert response.status_code == 403
    assert 'data-error="csrf"' in response.text


def test_import_is_recorded_in_the_usage_log(client, world):
    signed = sign_in(client, world["users"]["owner_a"]["email"])
    upload(client, signed, workbook_bytes([sample(CODE_A)]))
    usage = client.get("/usage")
    assert usage.status_code == 200
    assert "catalog.import" in usage.text
    assert "created=1" in usage.text


def test_csv_export_neutralises_spreadsheet_formulas(client, world):
    tenant_a = world["tenants"]["a"]
    conn = world["conn"]
    conn.execute(
        "INSERT INTO medicines (tenant_id, barcode, name, verified, created_at, updated_at) "
        "VALUES (?, ?, ?, 1, '2026-10-22T00:00:00+00:00', '2026-10-22T00:00:00+00:00')",
        (tenant_a, ean13("955001000904"), "=HYPERLINK(\"http://evil.example\")"),
    )
    medicine_id = conn.execute(
        "SELECT id FROM medicines WHERE tenant_id = ? AND name LIKE '=%'", (tenant_a,)
    ).fetchone()["id"]
    conn.execute(
        "INSERT INTO scans (tenant_id, user_id, barcode, medicine_id, source, found, created_at) "
        "VALUES (?, ?, ?, ?, 'camera', 1, '2026-10-22T00:00:00+00:00')",
        (tenant_a, world["users"]["owner_a"]["id"], ean13("955001000904"), medicine_id),
    )
    conn.commit()

    sign_in(client, world["users"]["owner_a"]["email"])
    export = client.get("/export/scans.csv")
    assert export.status_code == 200
    assert "'=HYPERLINK" in export.text
    assert "\n=HYPERLINK" not in export.text
